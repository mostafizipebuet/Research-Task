"""Strict input validation. Repeated explanations are observations, not duplicates to drop."""
from pathlib import Path
import hashlib
import json
import re
from collections import Counter
import openpyxl
import pandas as pd

class DataError(ValueError):
    pass

def sha256(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()

def unique_object(pairs):
    obj = {}
    for k, v in pairs:
        if k in obj:
            raise DataError(f'Duplicate JSON key: {k}')
        obj[k] = v
    return obj

def canonical_requirement(value):
    m = re.fullmatch(r'(FR|PR|ER|RR)\.?([0-9]+)', value)
    if not m:
        raise DataError(f'Invalid requirement ID: {value!r}')
    return f'{m[1]}.{int(m[2])}'

def parse_workbook(path, criteria, reference, expected_runs=100):
    req_ids = [r['requirement_id'] for r in reference['requirements']]
    names = {c['criterion_id']: c['name'] for c in criteria}
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=False)
    if len(workbook.sheetnames) != expected_runs:
        raise DataError(f'Expected {expected_runs} worksheets, found {len(workbook.sheetnames)}. No sheets were silently skipped.')
    flags, opportunities, runs, raw_hashes, warnings = [], [], [], [], []
    for sheet in workbook:
        if sheet['A1'].value != 'Raw Data':
            raise DataError(f'{sheet.title}: A1 must contain Raw Data')
        value = sheet['A2'].value
        if not isinstance(value, str) or sheet['A2'].data_type == 'f':
            raise DataError(f'{sheet.title}: A2 must contain literal JSON text')
        try:
            data = json.loads(value, object_pairs_hook=unique_object)
        except (json.JSONDecodeError, DataError) as exc:
            raise DataError(f'{sheet.title}: {exc}') from exc
        if not isinstance(data, dict):
            raise DataError(f'{sheet.title}: top-level JSON must be an object')
        mapped = {}
        for key in data:
            canonical = canonical_requirement(key)
            if canonical in mapped:
                raise DataError(f'{sheet.title}: duplicate normalized requirement {canonical}')
            mapped[canonical] = key
        if set(mapped) != set(req_ids):
            raise DataError(f'{sheet.title}: requirement IDs do not match the 13-item reference; missing={set(req_ids)-set(mapped)}, extra={set(mapped)-set(req_ids)}')
        runs.append(sheet.title)
        raw_hashes.append(hashlib.sha256(value.encode('utf-8')).hexdigest())
        for req in req_ids:
            source_key = mapped[req]
            items = data[source_key]
            if not isinstance(items, list):
                raise DataError(f'{sheet.title}/{req}: expected a list')
            seen = set()
            for index, item in enumerate(items):
                if not isinstance(item, dict) or not {'criterion','name','explanation'}.issubset(item):
                    raise DataError(f'{sheet.title}/{req}/{index}: incomplete flag object')
                crit = item['criterion']
                if crit not in names or item['name'] != names[crit]:
                    raise DataError(f'{sheet.title}/{req}: invalid criterion/name {crit}/{item["name"]}')
                if crit in seen:
                    raise DataError(f'{sheet.title}/{req}: duplicate criterion {crit}; resolve before counting')
                seen.add(crit)
                text = item['explanation']
                if not isinstance(text, str) or not text.strip():
                    raise DataError(f'{sheet.title}/{req}/{crit}: empty explanation')
                flags.append(dict(record_id=f'{sheet.title}|{req}|{crit}', run_id=sheet.title, requirement_id=req,
                                  source_requirement_key=source_key, criterion_id=crit, criterion_name=names[crit],
                                  source_cell='A2', json_pointer=f'/{source_key}/{index}', explanation=text,
                                  explanation_sha256=hashlib.sha256(text.encode('utf-8')).hexdigest()))
            for crit in names:
                opportunities.append(dict(run_id=sheet.title,requirement_id=req,criterion_id=crit,flagged=int(crit in seen)))
        if sheet.max_row > 2 or sheet.max_column > 1:
            warnings.append(f'{sheet.title}: worksheet has dimensions beyond A1:A2; analysis reads documented cell A2 only.')
    workbook.close()
    numeric_ids = [int(x.replace('Execution','')) for x in runs if re.fullmatch(r'Execution\d+', x)]
    gaps = sorted(set(range(min(numeric_ids), max(numeric_ids)+1))-set(numeric_ids)) if numeric_ids else []
    if gaps:
        warnings.append(f'Worksheet label gap(s): {gaps}. Actual worksheet count is {len(runs)}. Labels are identifiers, not time or proof of a missing run.')
    flag_df = pd.DataFrame(flags)
    opp_df = pd.DataFrame(opportunities)
    if int(opp_df.flagged.sum()) != len(flags):
        raise DataError('Flags do not reconcile with the complete opportunity table')
    if flag_df.record_id.duplicated().any():
        raise DataError('Record IDs are not unique')
    audit = dict(workbook_sha256=sha256(path), n_runs=len(runs), n_requirements=len(req_ids), n_criteria=len(names),
                 n_opportunities=len(opportunities), n_flags=len(flags), n_unflagged=len(opportunities)-len(flags),
                 n_unique_explanation_texts=int(flag_df.explanation.nunique()), n_unique_raw_run_json=len(set(raw_hashes)),
                 n_observed_requirement_criterion_pairs=int(flag_df.groupby(['requirement_id','criterion_id']).ngroups),
                 flags_by_criterion=dict(Counter(flag_df.criterion_id)), run_ids=runs, naming_gaps=gaps, warnings=warnings,
                 missing_criterion_interpretation='Not flagged by the LLM in that run; not a confirmed pass.',
                 independence='Stated by the task description. Not independently verified from these files.')
    return flag_df, opp_df, audit

def write_csv(frame, path):
    """UTF-8 BOM supports direct Excel opening; raw text is not rewritten."""
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    frame.to_csv(path, index=False, encoding='utf-8-sig', float_format='%.8g')
